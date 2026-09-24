from __future__ import annotations

import json
from pathlib import Path

import pdfplumber
from docx import Document
from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Sebastian Mercado")
OUT = Path(r"tmp\project-diagnosis\extracted")
OUT.mkdir(parents=True, exist_ok=True)

SOURCES = {
    "business_case": ROOT / r"Documents\UNIVERSIDAD\CAPSTONE PROJECT SISTEMAS\SEMANA 4\Business_Case_Capstone_V2.6.pdf",
    "project_charter": ROOT / r"Documents\UNIVERSIDAD\CAPSTONE PROJECT SISTEMAS\SEMANA 4\Acta_Constitucion - Capstone V2.pdf",
    "srs": ROOT / r"Documents\UNIVERSIDAD\CAPSTONE PROJECT SISTEMAS\SEMANA 4\Documento_Requerimientos_Torcly_V2.6.pdf",
    "analysis_structure": ROOT / r"Documents\UNIVERSIDAD\CAPSTONE PROJECT SISTEMAS\SEMANA 4\Estructura_de_Analisis_Torcly v2.2.pdf",
    "edt_diagram": ROOT / r"Documents\UNIVERSIDAD\CAPSTONE PROJECT SISTEMAS\SEMANA 3\Diagrama EDT.pdf",
    "product_backlog": ROOT / r"Downloads\Product_Backlog_Kanban_Torcly (1) (2).pdf",
    "daily": ROOT / r"Downloads\Daily_Kanban_Torcly_Completo (1).pdf",
    "development_plan": ROOT / r"Downloads\Plan_de_Desarrollo_Torcly_EDT_completa_horas (1).docx",
    "status_report": ROOT / r"Downloads\2_Informe de Estado del Proyecto (2).docx",
    "status_report_18_sep": ROOT / r"Downloads\Estado_del_Proyecto_Torcly_18-09-2026.docx",
    "status_report_prior": ROOT / r"Downloads\2_Informe de Estado del Proyecto.docx",
    "schedule": ROOT / r"Documents\UNIVERSIDAD\CAPSTONE PROJECT SISTEMAS\SEMANA 3\Cronograma - Capstone V2.5.xlsx",
    "costing": ROOT / r"Downloads\Costeo_Torcly_por_Paquete_de_Trabajo_V1.0.xlsx",
    "risks": ROOT / r"Downloads\Matriz_de_Riesgos_Torcly_corregida (1) (1).xlsx",
    "responsibilities": ROOT / r"Downloads\Matriz_de_Responsabilidades_Torcly_Actualizada_V2.5.xlsx",
    "communications": ROOT / r"Downloads\Matriz_de_Comunicaciones_Torcly_corregida.xlsx",
}


def clean(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def extract_pdf(key: str, path: Path) -> dict[str, object]:
    page_texts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            page_texts.append(f"=== PAGE {index} ===\n{page.extract_text() or ''}")
    (OUT / f"{key}.txt").write_text("\n\n".join(page_texts), encoding="utf-8")
    return {"type": "pdf", "path": str(path), "pages": len(page_texts)}


def extract_docx(key: str, path: Path) -> dict[str, object]:
    document = Document(path)
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(f"P: {text}")
    for table_index, table in enumerate(document.tables, start=1):
        blocks.append(f"\n=== TABLE {table_index} ===")
        for row in table.rows:
            blocks.append(" | ".join(clean(cell.text) for cell in row.cells))
    (OUT / f"{key}.txt").write_text("\n".join(blocks), encoding="utf-8")
    return {
        "type": "docx",
        "path": str(path),
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
    }


def extract_xlsx(key: str, path: Path) -> dict[str, object]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    output: list[str] = []
    sheets: list[dict[str, object]] = []
    for worksheet in workbook.worksheets:
        output.append(f"\n=== SHEET {worksheet.title} ({worksheet.max_row}x{worksheet.max_column}) ===")
        sheets.append({"name": worksheet.title, "rows": worksheet.max_row, "columns": worksheet.max_column})
        for row in worksheet.iter_rows():
            populated = [f"{cell.coordinate}={clean(cell.value)}" for cell in row if cell.value is not None]
            if populated:
                output.append(" | ".join(populated))
    (OUT / f"{key}.txt").write_text("\n".join(output), encoding="utf-8")
    return {"type": "xlsx", "path": str(path), "sheets": sheets}


manifest: dict[str, object] = {}
for source_key, source_path in SOURCES.items():
    if not source_path.exists():
        manifest[source_key] = {"path": str(source_path), "missing": True}
        continue
    suffix = source_path.suffix.lower()
    if suffix == ".pdf":
        manifest[source_key] = extract_pdf(source_key, source_path)
    elif suffix == ".docx":
        manifest[source_key] = extract_docx(source_key, source_path)
    elif suffix == ".xlsx":
        manifest[source_key] = extract_xlsx(source_key, source_path)

(OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(manifest, ensure_ascii=False, indent=2))
